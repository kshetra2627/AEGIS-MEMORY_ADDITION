"""Load PDF/DOCX/TXT files into page-level LangChain Documents with metadata."""
import os
import hashlib
from langchain_core.documents import Document


def file_hash(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _infer_title(pages_text: list[str], filename: str) -> str:
    for text in pages_text[:2]:
        for line in text.splitlines():
            line = line.strip()
            if 8 <= len(line) <= 120 and not line.isdigit():
                return line
    return os.path.splitext(filename)[0].replace("_", " ").replace("-", " ")


def load_pdf(path: str) -> list[Document]:
    from pypdf import PdfReader
    filename = os.path.basename(path)
    reader = PdfReader(path)
    pages_text = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:
            pages_text.append("")
    title = _infer_title(pages_text, filename)
    docs = []
    for i, text in enumerate(pages_text):
        if not text.strip():
            continue
        docs.append(Document(page_content=text, metadata={
            "filename": filename, "title": title, "page": i + 1,
        }))
    return docs


def load_docx(path: str) -> list[Document]:
    import docx
    filename = os.path.basename(path)
    d = docx.Document(path)
    full_text = "\n".join(p.text for p in d.paragraphs)
    title = _infer_title([full_text], filename)
    if not full_text.strip():
        return []
    return [Document(page_content=full_text, metadata={
        "filename": filename, "title": title, "page": 1,
    })]


def load_txt(path: str) -> list[Document]:
    filename = os.path.basename(path)
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    title = _infer_title([text], filename)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={
        "filename": filename, "title": title, "page": 1,
    })]


def load_file(path: str) -> list[Document]:
    """Loads a single file, returns [] and logs on corrupt/malformed input instead of raising."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return load_pdf(path)
        elif ext == ".docx":
            return load_docx(path)
        elif ext == ".txt":
            return load_txt(path)
        else:
            return []
    except Exception as e:
        print(f"[loader] Skipping corrupted/unsupported file {path}: {e}")
        return []
